"""Export Sprint 1 deliverables from pm/ Markdown sources.

Usage: python3 tools/export_s1.py  (needs pandoc, openpyxl)
Outputs to deliverables/S1/ with the HK261_SPM_FP_Team6_ prefix [03 §7].
"""
import re
import subprocess
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "deliverables" / "S1"
PREFIX = "HK261_SPM_FP_Team6_"


def pandoc(sources, target, title):
    cmd = ["pandoc", "--from", "gfm", "--toc", "--metadata", f"title={title}", "-o", str(target)]
    subprocess.run(cmd + [str(ROOT / s) for s in sources], check=True)


def md_tables(path):
    """Return every pipe table in a Markdown file as (heading, rows)."""
    tables, heading, current = [], "", None
    for line in (ROOT / path).read_text(encoding="utf-8").splitlines():
        if line.startswith("#"):
            heading = line.lstrip("# ").strip()
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", c) for c in cells if c):
                continue
            if current is None:
                current = (heading, [])
                tables.append(current)
            current[1].append([re.sub(r"\*\*|`", "", c) for c in cells])
        else:
            current = None
    return tables


HEAD = PatternFill("solid", fgColor="8A3B12")


def add_sheet(wb, name, path, pick=None, note=None):
    ws = wb.create_sheet(name)
    r = 1
    ws.cell(r, 1, f"Source: {path}").font = Font(italic=True, color="666666")
    r += 1
    if note:
        ws.cell(r, 1, note).font = Font(bold=True, color="B42318")
        r += 1
    for i, (heading, rows) in enumerate(md_tables(path)):
        if pick and i not in pick:
            continue
        r += 1
        ws.cell(r, 1, heading).font = Font(bold=True, size=12)
        r += 1
        for j, row in enumerate(rows):
            for k, val in enumerate(row, start=1):
                c = ws.cell(r, k, val)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if j == 0:
                    c.font = Font(bold=True, color="FFFFFF")
                    c.fill = HEAD
            r += 1
    for col in ws.columns:
        width = max((len(str(c.value or "")) for c in col[2:]), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max(width, 10), 60)
    return ws


def workbook(target):
    wb = Workbook()
    wb.remove(wb.active)
    plan_note = "Values marked TBD need team input (plan/INPUTS.md). Plan figures are not actuals."
    add_sheet(wb, "ProductBacklog", "pm/backlog/product-backlog.md", pick=[1, 2, 3, 4], note=plan_note)
    add_sheet(wb, "DoR", "pm/backlog/definition-of-ready.md")
    add_sheet(wb, "PlanningPoker", "pm/backlog/planning-poker-log.md",
              note="Votes are entered by the team only. Attach photos of each reveal.")
    add_sheet(wb, "ReleasePlan", "pm/pmp/03-master-release-plan.md", note=plan_note)
    add_sheet(wb, "S1_Backlog", "pm/sprints/S1/sprint-backlog.md", note=plan_note)
    burn = wb.create_sheet("Burndown")
    burn.append(["Day", "Date", "Ideal remaining (h, plan)", "Actual remaining (h, Jira)"])
    days = ["Fri 23/10", "Mon 26/10", "Tue 27/10", "Wed 28/10", "Thu 29/10"]
    total = 48
    for i, d in enumerate(days):
        burn.append([i + 1, d, round(total * (1 - (i + 1) / len(days)), 1), "TBD"])
    burn.append([])
    burn.append(["Ideal line uses the proposed 48 h; actual values must come from Jira Completed Work."])
    add_sheet(wb, "RiskRegister", "pm/pmp/05-risk-plan.md", pick=[1])
    add_sheet(wb, "IssueLog", "pm/pmp/05-risk-plan.md", pick=[2])
    add_sheet(wb, "ChangeLog", "pm/change/change-log.md")
    add_sheet(wb, "ConfigItems", "pm/pmp/06-configuration-management.md", pick=[1, 2])
    add_sheet(wb, "TestCases_F1", "pm/sprints/S1/test-cases-F1.md")
    wb.save(target)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    pmp = sorted(str(p.relative_to(ROOT)) for p in (ROOT / "pm" / "pmp").glob("0*.md"))
    pandoc(pmp, OUT / f"{PREFIX}c1_ProjectPlan.docx", "Project Management Plan v1.0 – Team 6 Page Turners")
    report = ["pm/sprints/S1/sprint-report.md", "pm/sprints/S1/sprint-backlog.md", "pm/sprints/S1/daily-log.md",
              "pm/sprints/S1/test-cases-F1.md", "pm/sprints/S1/review-retro.md"]
    pandoc(report, OUT / f"{PREFIX}c3_SprintReport_S1.docx", "Sprint 1 Report – Team 6 Page Turners")
    pandoc(["pm/pmp/03-master-release-plan.md"], OUT / f"{PREFIX}MasterPlan_Timeline.docx",
           "Master Plan and Timeline – Team 6 Page Turners")
    workbook(OUT / f"{PREFIX}c2_SprintBacklog_S1.xlsx")
    print("Exported:", *sorted(p.name for p in OUT.iterdir()), sep="\n  ")


if __name__ == "__main__":
    main()
